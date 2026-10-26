import os
from django.core.management.base import BaseCommand
from docx import Document

class Command(BaseCommand):
    help = 'сохраняем тесты джанго в ворд файл кароче'

    def handle(self, *args, **options):
        # создаем ворд документ
        doc = Document()
        doc.add_heading('Отчет по тестам Django', 0)

        doc.add_paragraph('Автоматический экспорт тестов из приложения blog.')

        # читаем тесты из файла tests.py
        tests_path = os.path.join('blog', 'tests.py')
        if os.path.exists(tests_path):
            with open(tests_path, 'r', encoding='utf-8') as f:
                code = f.read()

            doc.add_heading('Исходный код тестов (blog/tests.py)', level=1)
            # добавляем код в ворд
            doc.add_paragraph(code)
            
            # сохраняем файл
            output_file = 'django_tests_report.docx'
            doc.save(output_file)
            self.stdout.write(self.style.SUCCESS(f'все тесты записаны в файл {output_file}'))
        else:
            self.stdout.write(self.style.ERROR('файл с тестами не найден ваще'))